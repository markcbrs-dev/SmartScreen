import importlib.util
from io import BytesIO
from pathlib import Path
from unittest import skipUnless

from django.contrib.auth.models import User
from django.core import mail
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import connection
from django.test import TestCase
from django.test.utils import override_settings
from django.urls import reverse

from docx import Document

from main.models import Applicant, JobDescription
from main.services.recruitment import replace_active_job_description


def build_docx_file(name: str, paragraphs: list[str]) -> SimpleUploadedFile:
	buffer = BytesIO()
	document = Document()
	for paragraph in paragraphs:
		if isinstance(paragraph, dict):
			style = paragraph.get('style')
			doc_paragraph = document.add_paragraph(style=style)
			for run_text, is_bold in paragraph.get('runs', []):
				run = doc_paragraph.add_run(run_text)
				run.bold = bool(is_bold)
		elif isinstance(paragraph, tuple):
			text, style = paragraph
			document.add_paragraph(text, style=style)
		else:
			document.add_paragraph(paragraph)
	document.save(buffer)
	buffer.seek(0)
	return SimpleUploadedFile(
		name,
		buffer.getvalue(),
		content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
	)


def build_text_file(name: str, content: str) -> SimpleUploadedFile:
	return SimpleUploadedFile(name, content.encode('utf-8'), content_type='text/plain')


def build_workspace_file(name: str) -> SimpleUploadedFile:
	file_path = Path(__file__).resolve().parents[2] / name
	return SimpleUploadedFile(
		name,
		file_path.read_bytes(),
		content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
	)


class ApplicantFlowTests(TestCase):
	def test_home_shows_role_selection_links(self):
		response = self.client.get(reverse('home'))

		self.assertEqual(response.status_code, 200)
		self.assertContains(response, 'Continue as Applicant')
		self.assertContains(response, 'Continue as HR')
		self.assertContains(response, reverse('apply'))
		self.assertContains(response, reverse('login'))

	def test_apply_page_shows_uploaded_job_description_without_auto_match_option(self):
		job = replace_active_job_description(build_text_file(
			'computer-science-professor.txt',
			'Job Title: Computer Science Professor\nCompany: SmartScreen University\nLocation: Manila\nResponsibilities\nTeach computer science courses\nQualifications\nBachelor of Science in Computer Science\nPython Java JavaScript research teaching',
		))

		response = self.client.get(reverse('apply'))

		self.assertEqual(response.status_code, 200)
		self.assertContains(response, 'Select a target position')
		self.assertContains(response, 'Job Description')
		self.assertContains(response, job.title)
		self.assertNotContains(response, 'class="job-subtitle"', html=False)
		self.assertNotContains(response, 'novalidate')
		self.assertNotContains(response, 'Auto-match me to the best job')
		self.assertContains(response, 'id="applicantLayout" class="layout layout-collapsed"', html=False)
		self.assertContains(response, 'id="jobDescriptionCard" class="card job-card is-collapsed"', html=False)
		self.assertNotContains(response, 'Experience Level')
		self.assertNotContains(response, 'Required Skills')
		self.assertNotContains(response, '<h3 class="section-title">Qualifications</h3>', html=False)
		self.assertContains(response, 'Job Description')
		self.assertNotContains(response, 'Current Job Description')

	def test_apply_page_uses_about_the_job_text_for_job_description(self):
		replace_active_job_description(build_text_file(
			'faculty-role.txt',
			'Job Title: Computer Science Faculty\n'
			'Company: SmartScreen University\n'
			'Location: Manila\n'
			'About the job:\n'
			'Teach undergraduate computing courses.\n'
			'Guide student capstone projects.\n'
			'Job Description:\n'
			'Coordinate with faculty teams.\n',
		))

		response = self.client.post(
			reverse('apply'),
			{
				'job_id': '',
			},
		)

		self.assertEqual(response.status_code, 200)
		self.assertContains(response, 'About the job:')
		self.assertContains(response, 'Teach undergraduate computing courses.')
		self.assertContains(response, 'Guide student capstone projects.')
		self.assertContains(response, 'Job Description:')
		self.assertContains(response, 'Coordinate with faculty teams.')
		self.assertNotContains(response, '<h3 class="section-title">Responsibilities</h3>', html=False)
		self.assertContains(response, '<p class="job-description-copy">Teach undergraduate computing courses.</p>', html=False)
		self.assertContains(response, '<p class="job-description-copy">Guide student capstone projects.</p>', html=False)
		self.assertContains(response, '<p class="job-description-copy">Coordinate with faculty teams.</p>', html=False)

	def test_apply_page_keeps_about_the_job_bullets(self):
		replace_active_job_description(build_docx_file(
			'faculty-role.docx',
			[
				'Computer Science Faculty',
				'Company: SmartScreen University',
				'About the job',
				('Teach undergraduate IT courses.', 'List Bullet'),
				('Prepare lesson plans and assessments.', 'List Bullet'),
			],
		))

		response = self.client.get(reverse('apply'))

		self.assertEqual(response.status_code, 200)
		self.assertContains(response, '<p class="job-description-heading">About the job:</p>', html=False)
		self.assertContains(response, '<li>Teach undergraduate IT courses.</li>', html=False)
		self.assertContains(response, '<li>Prepare lesson plans and assessments.</li>', html=False)

	def test_apply_page_preserves_leading_dashes_in_job_description(self):
		replace_active_job_description(build_text_file(
			'faculty-role.txt',
			'Job Title: Computer Science Faculty\n'
			'About the job:\n'
			'- Can greet customers politely\n'
			'- Can code effectively\n',
		))

		response = self.client.get(reverse('apply'))

		self.assertEqual(response.status_code, 200)
		self.assertContains(response, '<li>- Can greet customers politely</li>', html=False)
		self.assertContains(response, '<li>- Can code effectively</li>', html=False)

	def test_apply_job_details_endpoint_returns_current_uploaded_jobs(self):
		first_job = replace_active_job_description(build_text_file(
			'front-desk-associate.txt',
			'Job Title: Front Desk Associate\n'
			'Company: SmartScreen Suites\n'
			'About the job:\n'
			'Welcome guests and manage reservations.\n',
		))
		replace_active_job_description(build_text_file(
			'barista.txt',
			'Job Title: Barista\n'
			'Company: SmartScreen Cafe\n'
			'Duties and Responsibilities:\n'
			'- Prepare beverages and assist customers.\n',
		))

		response = self.client.get(reverse('applicant_job_details'))

		self.assertEqual(response.status_code, 200)
		payload = response.json()
		self.assertTrue(payload['has_active_job'])
		self.assertEqual(len(payload['jobs']), 2)
		self.assertEqual(payload['jobs'][0]['title'], 'Barista')
		self.assertEqual(payload['jobs'][1]['title'], first_job.title)
		self.assertIn('Duties and Responsibilities:', [section['heading_html'] for section in payload['jobs'][0]['description_preview_sections']])

	def test_apply_matches_selected_job_and_stores_resume_analysis(self):
		job = replace_active_job_description(build_text_file(
			'computer-science-professor.txt',
			'Job Title: Computer Science Professor\nCompany: SmartScreen University\nLocation: Manila\nResponsibilities\nTeach computer science courses\nQualifications\nBachelor of Science in Computer Science\nPython Java JavaScript Artificial Intelligence Data Science teaching research',
		))

		response = self.client.post(
			reverse('apply'),
			{
				'job_id': str(job.id),
				'resume': build_docx_file(
					'resume.docx',
					[
						'Jane Applicant',
						'jane@example.com',
						'09123456789',
						'Python Java JavaScript teaching and research experience',
						'Bachelor of Science in Computer Science',
						'Artificial Intelligence and Data Science projects',
					],
				),
			},
			follow=True,
		)

		self.assertEqual(response.status_code, 200)
		self.assertContains(response, 'Application submitted successfully.')
		self.assertNotContains(response, 'Best match:')
		applicant = Applicant.objects.get()
		self.assertEqual(applicant.first_name, 'Jane')
		self.assertEqual(applicant.middle_name, '')
		self.assertEqual(applicant.last_name, 'Applicant')
		self.assertEqual(applicant.contact_number, '+63 912 345 6789')
		self.assertEqual(applicant.email, 'jane@example.com')
		self.assertEqual(applicant.position, job.title)
		self.assertEqual(applicant.applied_job, job)
		self.assertTrue(applicant.job_preference_locked)
		self.assertIn('Python', applicant.parsed_skills)
		self.assertGreaterEqual(applicant.match_score, 1)
		self.assertEqual(applicant.status, 'INTERVIEW')
		self.assertEqual(applicant.match_notes.get('matched_job', {}).get('job_title'), job.title)
		self.assertEqual(applicant.match_notes.get('scored_job_count'), 1)

	def test_apply_stores_middle_initial_separately_from_last_name(self):
		job = replace_active_job_description(build_text_file(
			'painter.txt',
			'Job Title: Painter\nCompany: SmartScreen Builders\nLocation: Antipolo\nPainter Putty Exterior painting',
		))

		response = self.client.post(
			reverse('apply'),
			{
				'job_id': str(job.id),
				'resume': build_docx_file(
					'marlou-resume.docx',
					[
						'Marlou R. Saysay',
						'marlousaysay99@gmail.com',
						'09451386454',
						'Painter Putty Exterior painting',
					],
				),
			},
			follow=True,
		)

		self.assertEqual(response.status_code, 200)
		applicant = Applicant.objects.get(email='marlousaysay99@gmail.com')
		self.assertEqual(applicant.first_name, 'Marlou')
		self.assertEqual(applicant.middle_name, 'R.')
		self.assertEqual(applicant.last_name, 'Saysay')

	def test_apply_keeps_status_new_when_match_score_is_below_70(self):
		job = replace_active_job_description(build_text_file(
			'backend-engineer.txt',
			'Job Title: Backend Engineer\nCompany: SmartScreen Labs\nLocation: Remote\nPython Django APIs Kubernetes',
		))

		response = self.client.post(
			reverse('apply'),
			{
				'job_id': str(job.id),
				'resume': build_docx_file(
					'low-match-resume.docx',
					[
						'Low Match Applicant',
						'lowmatch@example.com',
						'09123451111',
						'Customer service communication',
					],
				),
			},
			follow=True,
		)

		self.assertEqual(response.status_code, 200)
		applicant = Applicant.objects.get(email='lowmatch@example.com')
		self.assertLess(applicant.match_score, 70)
		self.assertEqual(applicant.status, 'NEW')

	def test_apply_scores_resume_against_all_uploaded_jobs_for_recommendation(self):
		selected_job = replace_active_job_description(build_text_file(
			'room-attendant.txt',
			'Job Title: Room Attendant\n'
			'Company: SmartScreen Resort\n'
			'Responsibilities\n'
			'Clean guest rooms and replace linens\n'
			'Qualifications\n'
			'Housekeeping Laundry Cleaning Bed Making\n',
		))
		recommended_job = replace_active_job_description(build_text_file(
			'kids-attendant.txt',
			'Job Title: Kids Attendant\n'
			'Company: SmartScreen Resort\n'
			'Responsibilities\n'
			'Supervise children and facilitate play activities\n'
			'Qualifications\n'
			'Childcare Play Activities Guest Engagement Communication First Aid\n',
		))

		response = self.client.post(
			reverse('apply'),
			{
				'job_id': str(selected_job.id),
				'resume': build_docx_file(
					'kids-attendant-resume.docx',
					[
						'Jamie Applicant',
						'jamie@example.com',
						'09123450000',
						'Childcare play activities guest engagement communication first aid',
						'Experienced in supervising children at family resort programs',
					],
				),
			},
			follow=True,
		)

		self.assertEqual(response.status_code, 200)
		applicant = Applicant.objects.get()
		self.assertEqual(applicant.applied_job, selected_job)
		self.assertEqual(applicant.position, selected_job.title)
		self.assertEqual(applicant.match_notes.get('scored_job_count'), 2)
		self.assertEqual(applicant.match_notes.get('recommendation', {}).get('job_title'), recommended_job.title)
		self.assertGreater(
			applicant.match_notes.get('recommendation', {}).get('overall_score', 0),
			applicant.match_notes.get('overall_score', 0),
		)

	@skipUnless(importlib.util.find_spec('en_core_web_sm') is not None, 'en_core_web_sm is not installed')
	def test_apply_with_real_spacy_ner_extracts_resume_location_from_docx(self):
		job = replace_active_job_description(build_text_file(
			'platform-engineer.txt',
			'Job Title: Platform Engineer\nCompany: SmartScreen Labs\nLocation: Seattle\nPython Docker Kubernetes cloud infrastructure',
		))

		response = self.client.post(
			reverse('apply'),
			{
				'job_id': str(job.id),
				'resume': build_docx_file(
					'barack-obama-resume.docx',
					[
						'Barack Obama',
						'barack@example.com',
						'09123456780',
						'Professional Experience',
						'Worked at Microsoft building cloud platforms in Seattle.',
						'Python Docker Kubernetes cloud infrastructure',
					],
				),
			},
			follow=True,
		)

		self.assertEqual(response.status_code, 200)
		self.assertContains(response, 'Application submitted successfully.')
		applicant = Applicant.objects.get()
		self.assertEqual(applicant.first_name, 'Barack')
		self.assertEqual(applicant.middle_name, '')
		self.assertEqual(applicant.last_name, 'Obama')
		self.assertEqual(applicant.contact_number, '+63 912 345 6780')
		self.assertEqual(applicant.email, 'barack@example.com')
		self.assertEqual(applicant.applied_job, job)
		self.assertIn('Seattle', applicant.parsed_locations)
		self.assertIn('Python', applicant.parsed_skills)
		self.assertGreater(applicant.match_score, 0)

	def test_apply_page_hides_manual_identity_fields(self):
		response = self.client.get(reverse('apply'))

		self.assertEqual(response.status_code, 200)
		self.assertNotContains(response, 'label for="contact"', html=False)
		self.assertNotContains(response, 'name="first_name"', html=False)
		self.assertNotContains(response, 'name="last_name"', html=False)
		self.assertNotContains(response, 'name="contact"', html=False)
		self.assertNotContains(response, 'name="email"', html=False)

	def test_apply_requires_name_contact_and_email_from_resume(self):
		job = replace_active_job_description(build_text_file(
			'computer-science-professor.txt',
			'Job Title: Computer Science Professor\nCompany: SmartScreen University\nLocation: Manila\nResponsibilities\nTeach computer science courses\nQualifications\nBachelor of Science in Computer Science\nPython Java JavaScript Artificial Intelligence Data Science teaching research',
		))

		response = self.client.post(
			reverse('apply'),
			{
				'job_id': str(job.id),
				'resume': build_docx_file(
					'resume-missing-contact.docx',
					[
						'Jane Applicant',
						'Python Java JavaScript teaching and research experience',
						'Bachelor of Science in Computer Science',
					],
				),
			},
		)

		self.assertEqual(response.status_code, 200)
		self.assertContains(response, 'We could not extract the following from the uploaded resume: contact number, email.')
		self.assertFalse(Applicant.objects.exists())


class JobUploadTests(TestCase):
	def setUp(self):
		self.user = User.objects.create_user(username='hr@example.com', password='secret123', is_staff=True)
		self.client.force_login(self.user)

	def test_dashboard_uploads_multiple_job_descriptions_from_one_file(self):
		response = self.client.post(
			reverse('dashboard'),
			{
				'action': 'upload_job_description',
				'job_description_file': build_text_file(
					'hotel-roles.txt',
					'JOB DESCRIPTION\n'
					'JOB TITLE\n'
					'KIDS ATTENDANT\n'
					'DUTIES & RESPONSIBILITIES\n'
					'• Supervising children in their own activities.\n\n'
					'JOB DESCRIPTION\n'
					'JOB TITLE\n'
					'BARTENDER\n'
					'DUTIES & RESPONSIBILITIES\n'
					'• Preparing beverages for guests.\n',
				),
			},
			follow=True,
		)

		self.assertEqual(response.status_code, 200)
		self.assertContains(response, 'Uploaded 2 job descriptions')
		self.assertEqual(JobDescription.objects.count(), 2)
		self.assertContains(response, 'Kids Attendant')
		self.assertContains(response, 'Bartender')

	def test_dashboard_uploads_multiple_job_description_files(self):
		response = self.client.post(
			reverse('dashboard'),
			{
				'action': 'upload_job_description',
				'job_description_file': [
					build_text_file(
						'backend-engineer.txt',
						'Job Title: Backend Engineer\nCompany: SmartScreen Labs\nLocation: Remote\nEmployment Type: Full Time\nPython Django APIs',
					),
					build_text_file(
						'frontend-engineer.txt',
						'Job Title: Frontend Engineer\nCompany: SmartScreen Labs\nLocation: Remote\nEmployment Type: Full Time\nJavaScript React CSS',
					),
				],
			},
			follow=True,
		)

		self.assertEqual(response.status_code, 200)
		self.assertContains(response, 'Uploaded 2 job descriptions')
		self.assertEqual(JobDescription.objects.count(), 2)
		self.assertTrue(JobDescription.objects.filter(title='Backend Engineer').exists())
		self.assertTrue(JobDescription.objects.filter(title='Frontend Engineer').exists())

	def test_dashboard_uploads_job_description(self):
		response = self.client.post(
			reverse('dashboard'),
			{
				'action': 'upload_job_description',
				'job_description_file': build_text_file(
					'computer-science-professor.txt',
					'Job Title: Computer Science Professor\nCompany: SmartScreen University\nLocation: Manila\nEmployment Type: Full Time\nPython Java JavaScript teaching research',
				),
			},
			follow=True,
		)

		self.assertEqual(response.status_code, 200)
		self.assertContains(response, 'Uploaded Job Description')

	def test_dashboard_uploads_real_multi_jd_docx_file(self):
		response = self.client.post(
			reverse('dashboard'),
			{
				'action': 'upload_job_description',
				'job_description_file': build_workspace_file('JOB-DESCRIPTION-all-works.docx'),
			},
			follow=True,
		)

		self.assertEqual(response.status_code, 200)
		self.assertContains(response, 'Uploaded ')
		self.assertGreater(JobDescription.objects.count(), 1)
		self.assertTrue(JobDescription.objects.filter(title='Butler').exists())
		self.assertTrue(JobDescription.objects.filter(title='Kids Attendant').exists())
		butler = JobDescription.objects.get(title='Butler')
		self.assertEqual(butler.reporting_to, 'Area Supervisor')
		self.assertIn('Hospitality', butler.industries)
		self.assertIn('Guest Service', butler.required_skills)

	def test_dashboard_keeps_existing_job_cards_when_uploading_new_job(self):
		self.client.post(
			reverse('dashboard'),
			{
				'action': 'upload_job_description',
				'job_description_file': build_text_file(
					'computer-science-professor.txt',
					'Job Title: Computer Science Professor\nCompany: SmartScreen University\nLocation: Manila\nPython Java JavaScript teaching research',
				),
			},
		)

		response = self.client.post(
			reverse('dashboard'),
			{
				'action': 'upload_job_description',
				'job_description_file': build_text_file(
					'backend-engineer.txt',
					'Job Title: Backend Engineer\nCompany: SmartScreen Labs\nLocation: Remote\nPython Django APIs',
				),
			},
			follow=True,
		)

		self.assertEqual(response.status_code, 200)
		self.assertEqual(JobDescription.objects.count(), 2)
		self.assertContains(response, 'Computer Science Professor')
		self.assertContains(response, 'Backend Engineer')

	def test_dashboard_prefers_job_description_section_over_qualifications_block(self):
		response = self.client.post(
			reverse('dashboard'),
			{
				'action': 'upload_job_description',
				'job_description_file': build_text_file(
					'faculty-role.txt',
					'Job Title: Computer Science/IT Faculty\n'
					'Company: FEU Institute of Technology\n'
					'Qualifications:\n'
					'Master\'s Degree in Computer Science or Information Technology\n'
					'At least three (3) years of CS/IT work\n'
					'Job Description:\n'
					'Develop and deliver curriculum in Information Technology subjects.\n'
					'Conduct lectures, labs, and seminars.\n',
				),
			},
			follow=True,
		)

		self.assertEqual(response.status_code, 200)
		self.assertContains(response, 'Job Description')
		self.assertContains(response, 'Develop and deliver curriculum in Information Technology subjects.')
		self.assertContains(response, 'Conduct lectures, labs, and seminars.')
		self.assertContains(response, '<p class="db-job-description-heading">Job Description:</p>', html=False)
		self.assertContains(response, '<p class="db-job-description-copy">Develop and deliver curriculum in Information Technology subjects.</p>', html=False)
		self.assertContains(response, '<p class="db-job-description-copy">Conduct lectures, labs, and seminars.</p>', html=False)

	def test_dashboard_preview_uses_about_the_job_section_and_keeps_docx_bullets(self):
		response = self.client.post(
			reverse('dashboard'),
			{
				'action': 'upload_job_description',
				'job_description_file': build_docx_file(
					'faculty-role.docx',
					[
						'Computer Science/IT Faculty',
						'Company: FEU Institute of Technology',
						'Location: Manila',
						'About the job',
						('Teach undergraduate IT courses.', 'List Bullet'),
						('Prepare lesson plans and assessments.', 'List Bullet'),
						'Job Description:',
						'Develop and deliver curriculum in Information Technology subjects.',
					],
				),
			},
			follow=True,
		)

		self.assertEqual(response.status_code, 200)
		self.assertContains(response, 'Description Preview')
		self.assertContains(response, '<p class="db-job-description-heading">About the job:</p>', html=False)
		self.assertContains(response, 'Teach undergraduate IT courses.')
		self.assertContains(response, 'Prepare lesson plans and assessments.')
		self.assertContains(response, '<ul class="db-job-detail-list">', html=False)
		self.assertContains(response, '<li>Teach undergraduate IT courses.</li>', html=False)
		self.assertContains(response, '<li>Prepare lesson plans and assessments.</li>', html=False)

	def test_dashboard_preserves_docx_bold_text_in_description_preview(self):
		response = self.client.post(
			reverse('dashboard'),
			{
				'action': 'upload_job_description',
				'job_description_file': build_docx_file(
					'faculty-role.docx',
					[
						'Computer Science/IT Faculty',
						'About the job',
						{'style': 'List Bullet', 'runs': [('Teach ', False), ('Python', True), (' and AI courses.', False)]},
					],
				),
			},
			follow=True,
		)

		self.assertEqual(response.status_code, 200)
		self.assertContains(response, '<li>Teach <strong>Python</strong> and AI courses.</li>', html=False)

	def test_dashboard_preserves_docx_bold_only_lines_in_description_preview(self):
		response = self.client.post(
			reverse('dashboard'),
			{
				'action': 'upload_job_description',
				'job_description_file': build_docx_file(
					'Python Programming Instructor.docx',
					[
						'About the job',
						{'runs': [('Responsibilities', True)]},
						{'runs': [('Curriculum and Instruction Development', True)]},
						('* Conduct classroom instruction and lesson planning aligned with the existing syllabus.', 'Normal'),
					],
				),
			},
			follow=True,
		)

		self.assertEqual(response.status_code, 200)
		self.assertContains(response, '<p class="db-job-description-heading"><strong>Responsibilities</strong>:</p>', html=False)
		self.assertContains(response, '<p class="db-job-description-copy"><strong>Curriculum and Instruction Development</strong></p>', html=False)

	def test_dashboard_dialog_uses_about_the_job_section_with_formatting(self):
		response = self.client.post(
			reverse('dashboard'),
			{
				'action': 'upload_job_description',
				'job_description_file': build_text_file(
					'faculty-role.txt',
					'Job Title: Computer Science/IT Faculty\n'
					'Overview:\n'
					'- Support **laboratory** classes.\n'
					'About the job:\n'
					'- Teach **undergraduate** IT courses.\n'
					'- Prepare lesson plans.\n',
				),
			},
			follow=True,
		)

		self.assertEqual(response.status_code, 200)
		self.assertNotContains(response, '<p class="db-job-description-heading">Overview:</p>', html=False)
		self.assertNotContains(response, 'Support <strong>laboratory</strong> classes.', html=False)
		self.assertContains(response, '<p class="db-job-description-heading">About the job:</p>', html=False)
		self.assertContains(response, '<p class="db-job-description-copy">Teach <strong>undergraduate</strong> IT courses.</p>', html=False)
		self.assertContains(response, '<p class="db-job-description-copy">Prepare lesson plans.</p>', html=False)

	def test_dashboard_prefers_filename_title_when_docx_body_contains_misleading_role_line(self):
		response = self.client.post(
			reverse('dashboard'),
			{
				'action': 'upload_job_description',
				'job_description_file': build_docx_file(
					'Python Programming Instructor.docx',
					[
						'About the job',
						'Collaborate with the Specialist',
						'Teach Python programming fundamentals.',
					],
				),
			},
			follow=True,
		)

		self.assertEqual(response.status_code, 200)
		self.assertContains(response, 'Uploaded job description: Python Programming Instructor.')
		self.assertContains(response, '<h3 class="db-job-detail-name">Python Programming Instructor</h3>', html=False)
		self.assertNotContains(response, '<h3 class="db-job-detail-name">Collaborate with the Specialist</h3>', html=False)

	def test_dashboard_limits_description_preview_to_250_characters(self):
		response = self.client.post(
			reverse('dashboard'),
			{
				'action': 'upload_job_description',
				'job_description_file': build_text_file(
					'faculty-role.txt',
					'Job Title: Computer Science/IT Faculty\n'
					'About the job:\n'
					'AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA\n'
					'BBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBB\n'
					'CCCCCCCCCCCCCCCCCCCCCCCCCCCCCCCCCCCCCCCCCCCCCCCCCCCCCCCCCCCCCCCCCCCCCCCCCCCCCCCCCCCCCCCCCCCCCCCCCCCC\n'
					'DDDDDDDDDDDDDDDDDDDDDDDDDDDDDDDDDDDDDDDDDDDDDDDDDDDDDDDDDDDDDDDDDDDDDDDDDDDDDDDDDD\n',
				),
			},
			follow=True,
		)

		self.assertEqual(response.status_code, 200)
		self.assertContains(response, 'AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA')
		self.assertContains(response, 'BBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBBB')
		self.assertContains(response, '<p class="db-job-description-copy">CCCCCCCCCCCCCCCCCCCCCCCCCCCCCCCCC...</p>', html=False)

	def test_dashboard_refreshes_existing_applicant_match_scores(self):
		replace_active_job_description(build_text_file(
			'computer-science-professor.txt',
			'Job Title: Computer Science Professor\nCompany: SmartScreen University\nLocation: Manila\nPython Java JavaScript Artificial Intelligence Data Science teaching research',
		))

		Applicant.objects.create(
			first_name='Mark',
			last_name='Cabreros',
			contact_number='+63 912 345 6789',
			email='mark@example.com',
			resume=build_docx_file(
				'mark-resume.docx',
				[
					'Mark Cabreros',
					'mark@example.com',
					'+63 912 345 6789',
					'Python Java JavaScript Artificial Intelligence Data Science',
					'Bachelor of Science in Computer Science',
					'Teaching and research experience',
				],
			),
		)

		response = self.client.get(reverse('dashboard'))

		self.assertEqual(response.status_code, 200)
		applicant = Applicant.objects.get()
		self.assertEqual(applicant.position, 'Computer Science Professor')
		self.assertIsNotNone(applicant.applied_job)
		self.assertGreater(applicant.match_score, 0)
		self.assertTrue(applicant.match_notes)
		self.assertEqual(
			applicant.match_notes.get('overall_score'),
			sum(applicant.match_notes.get('breakdown', {}).values()),
		)

	def test_dashboard_repairs_previously_saved_bad_applicant_name(self):
		job = replace_active_job_description(build_text_file(
			'python-instructor.txt',
			'Job Title: Python Programming Instructor\nCompany: FEU Institute of Technology\nLocation: Manila\nPython JavaScript teaching CSS',
		))

		Applicant.objects.create(
			first_name='Feu',
			last_name='Roosevelt',
			contact_number='+63 912 333 4455',
			email='faelman.jeromelemuel@gmail.com',
			resume=SimpleUploadedFile(
				'Jerome Lemuel Faelman - Resume.pdf',
				b'%PDF-1.4\n%',
				content_type='application/pdf',
			),
			applied_job=job,
			job_preference_locked=True,
			position=job.title,
			match_score=20,
			match_notes={
				'overall_score': 20,
				'breakdown': {
					'skill_score': 10,
					'experience_score': 5,
					'education_score': 5,
				},
			},
		)

		response = self.client.get(reverse('dashboard'))

		self.assertEqual(response.status_code, 200)
		applicant = Applicant.objects.get(email='faelman.jeromelemuel@gmail.com')
		self.assertEqual(applicant.first_name, 'Jerome')
		self.assertEqual(applicant.middle_name, 'Lemuel')
		self.assertEqual(applicant.last_name, 'Faelman')

	def test_dashboard_recalculates_stale_breakdown_format(self):
		job = replace_active_job_description(build_text_file(
			'backend-role.txt',
			'Job Title: Backend Engineer\nCompany: SmartScreen Labs\nPython Django APIs\nBachelor of Science in Computer Science',
		))

		Applicant.objects.create(
			first_name='Legacy',
			last_name='Applicant',
			contact_number='+63 912 111 2233',
			email='legacy@example.com',
			resume=build_docx_file(
				'legacy.docx',
				[
					'Legacy Applicant',
					'legacy@example.com',
					'+63 912 111 2233',
					'Python Django',
					'Bachelor of Science in Computer Science',
				],
			),
			applied_job=job,
			position=job.title,
			match_score=20,
			match_notes={
				'overall_score': 20,
				'breakdown': {
					'skill_score': 0,
					'experience_score': 50,
					'education_score': 50,
				},
			},
		)

		response = self.client.get(reverse('dashboard'))

		self.assertEqual(response.status_code, 200)
		applicant = Applicant.objects.get(email='legacy@example.com')
		self.assertEqual(
			applicant.match_notes.get('overall_score'),
			sum(applicant.match_notes.get('breakdown', {}).values()),
		)

	def test_dashboard_reparses_stale_degree_domain_job_skills(self):
		job = JobDescription.objects.create(
			title='Computer/IT Instructor',
			company='FEU Institute of Technology',
			location='Manila',
			employment_type='Full-Time or Part-Time',
			required_skills=['Information Technology', 'Computer Science'],
			preferred_skills=['CSS', 'Programming', 'Teaching'],
			education_requirements=['With Master\'s degree in IT/CS or any IT-related programs is a plus'],
			raw_text=(
				'Computer/IT Instructor\n'
				'Company: FEU Institute of Technology\n'
				'Location: Sampaloc, Manila\n'
				'Employment Type: Full-Time or Part-Time\n'
				'Qualifications\n'
				'• Graduate of BSIT/BSCS.\n'
				'• With Master\'s degree in IT/CS or any IT-related programs is a plus.\n'
				'• A holder of applicable National Certification in CSS and Programming is a plus.\n'
			),
		)

		Applicant.objects.create(
			first_name='Jerome',
			last_name='Lemuel',
			contact_number='+63 912 333 4455',
			email='jerome@example.com',
			resume=build_docx_file(
				'jerome.docx',
				[
					'Jerome Lemuel',
					'jerome@example.com',
					'+63 912 333 4455',
					'CSS Programming',
					'Bachelor of Science in Information Technology',
				],
			),
			applied_job=job,
			position=job.title,
			parsed_skills=['CSS', 'Programming'],
			parsed_education=['Bachelor of Science in Information Technology'],
			match_notes={
				'overall_score': 0,
				'matching_skills': [],
				'missing_skills': ['Information Technology', 'Computer Science'],
				'breakdown': {'skill_score': 0},
			},
		)

		response = self.client.get(reverse('dashboard'))

		self.assertEqual(response.status_code, 200)
		job.refresh_from_db()
		self.assertNotIn('Information Technology', job.required_skills)
		self.assertNotIn('Computer Science', job.required_skills)
		self.assertIn('CSS', job.preferred_skills)
		self.assertIn('Programming', job.preferred_skills)
		applicant = Applicant.objects.get(email='jerome@example.com')
		self.assertNotIn('Information Technology', applicant.match_notes.get('missing_skills', []))
		self.assertNotIn('Computer Science', applicant.match_notes.get('missing_skills', []))

	def test_dashboard_keeps_applicant_matched_to_selected_job(self):
		selected_job = replace_active_job_description(build_text_file(
			'computer-science-professor.txt',
			'Job Title: Computer Science Professor\nCompany: SmartScreen University\nLocation: Manila\nPython Java JavaScript teaching research',
		))
		replace_active_job_description(build_text_file(
			'backend-engineer.txt',
			'Job Title: Backend Engineer\nCompany: SmartScreen Labs\nLocation: Remote\nPython Django APIs microservices',
		))

		Applicant.objects.create(
			first_name='Jane',
			last_name='Applicant',
			contact_number='+63 912 345 6789',
			email='jane-selected@example.com',
			resume=build_docx_file(
				'jane-selected.docx',
				[
					'Jane Applicant',
					'jane-selected@example.com',
					'+63 912 345 6789',
					'Python Java JavaScript teaching research',
					'Bachelor of Science in Computer Science',
				],
			),
			applied_job=selected_job,
			job_preference_locked=True,
			position=selected_job.title,
		)

		response = self.client.get(reverse('dashboard'))

		self.assertEqual(response.status_code, 200)
		applicant = Applicant.objects.get()
		self.assertEqual(applicant.applied_job, selected_job)
		self.assertEqual(applicant.position, selected_job.title)

	def test_dashboard_removes_uploaded_job_description(self):
		job = replace_active_job_description(build_text_file(
			'backend-role.txt',
			'Job Title: Backend Engineer\nCompany: SmartScreen Labs\nLocation: Remote\nPython Django APIs',
		))

		response = self.client.post(
			reverse('dashboard'),
			{'action': 'remove_job_description', 'job_id': str(job.id)},
			follow=True,
		)

		self.assertEqual(response.status_code, 200)
		self.assertEqual(JobDescription.objects.count(), 0)
		self.assertContains(response, 'Uploaded job description removed.')

	def test_dashboard_updates_uploaded_job_description_from_edit_action(self):
		job = replace_active_job_description(build_text_file(
			'backend-role.txt',
			'Job Title: Backend Engineer\nCompany: SmartScreen Labs\nLocation: Remote\nPython Django APIs',
		))
		Applicant.objects.create(
			first_name='Jamie',
			last_name='Applicant',
			contact_number='+63 912 345 0000',
			email='jamie@example.com',
			resume=build_docx_file(
				'jamie.docx',
				[
					'Jamie Applicant',
					'jamie@example.com',
					'+63 912 345 0000',
				],
			),
			applied_job=job,
			position=job.title,
		)

		response = self.client.post(
			reverse('dashboard'),
			{
				'action': 'update_job_description',
				'job_id': str(job.id),
				'title': 'Senior Backend Engineer',
				'company': 'SmartScreen Labs PH',
				'location': 'Manila',
				'employment_type': 'Full-Time',
				'reporting_to': 'Engineering Manager',
				'raw_text': 'Job Title: Senior Backend Engineer\nCompany: SmartScreen Labs PH\nLocation: Manila\nReporting To: Engineering Manager\nAbout the job:\nBuild resilient APIs and mentor backend developers.\nJob Requirements:\n- Python Django APIs\n- Can mentor backend developers',
			},
			follow=True,
		)

		self.assertEqual(response.status_code, 200)
		job.refresh_from_db()
		self.assertEqual(job.title, 'Senior Backend Engineer')
		self.assertEqual(job.company, 'SmartScreen Labs PH')
		self.assertEqual(job.location, 'Manila')
		self.assertEqual(job.employment_type, 'Full-Time')
		self.assertEqual(job.reporting_to, 'Engineering Manager')
		self.assertIn('Build resilient APIs and mentor backend developers.', job.raw_text)
		self.assertContains(response, '<p class="db-job-description-heading">Job Requirements:</p>', html=False)
		self.assertContains(response, '<p class="db-job-description-copy">- Python Django APIs</p>', html=False)
		self.assertContains(response, 'Build resilient APIs and mentor backend developers.')
		self.assertContains(response, 'Updated job description: Senior Backend Engineer.')
		self.assertContains(response, 'data-dialog-id="jobEditDialog-{}'.format(job.id), html=False)
		applicant = Applicant.objects.get()
		self.assertEqual(applicant.position, 'Senior Backend Engineer')

	def test_dashboard_renders_contact_dialog_actions_in_status_tab(self):
		Applicant.objects.create(
			first_name='Avery',
			last_name='Santos',
			contact_number='+63 912 000 1122',
			email='avery@example.com',
			position='Backend Engineer',
			status='INTERVIEW',
			resume=build_docx_file(
				'avery.docx',
				[
					'Avery Santos',
					'avery@example.com',
					'+63 912 000 1122',
				],
			),
		)

		response = self.client.get(reverse('dashboard'))

		self.assertEqual(response.status_code, 200)
		self.assertContains(response, 'id="contactDialog"', html=False)
		self.assertContains(response, 'class="db-btn-contact js-open-contact-dialog"', html=False)
		self.assertContains(response, 'data-contact-applicant-id="', html=False)
		self.assertContains(response, 'data-contact-email="avery@example.com"', html=False)
		self.assertContains(response, 'data-contact-phone="+63 912 000 1122"', html=False)
		self.assertContains(response, 'id="contactDialogEmailAction"', html=False)
		self.assertContains(response, 'id="contactDialogSmsAction"', html=False)
		self.assertContains(response, 'https://mail.google.com/mail/?', html=False)
		self.assertContains(response, 'Send Email', html=False)
		self.assertContains(response, 'Send SMS', html=False)

	def test_dashboard_refresh_promotes_high_match_applicant_to_interview(self):
		job = replace_active_job_description(build_text_file(
			'backend-role.txt',
			'Job Title: Backend Engineer\nCompany: SmartScreen Labs\nLocation: Remote\nPython Django APIs',
		))
		applicant = Applicant.objects.create(
			first_name='Jordan',
			last_name='Cruz',
			contact_number='+63 912 444 7788',
			email='jordan@example.com',
			resume=build_docx_file(
				'jordan.docx',
				[
					'Jordan Cruz',
					'jordan@example.com',
					'+63 912 444 7788',
					'Python Django APIs',
				],
			),
			applied_job=job,
			position=job.title,
			status='NEW',
			parsed_skills=['Python', 'Django', 'REST API'],
			parsed_education=['Bachelor of Science in Information Technology'],
			match_notes={},
		)

		response = self.client.get(reverse('dashboard'))

		self.assertEqual(response.status_code, 200)
		applicant.refresh_from_db()
		self.assertGreaterEqual(applicant.match_score, 70)
		self.assertEqual(applicant.status, 'INTERVIEW')

	def test_dashboard_renders_score_analysis_dialog_for_scored_applicants(self):
		job = replace_active_job_description(build_text_file(
			'backend-role.txt',
			'Job Title: Backend Engineer\nCompany: SmartScreen Labs\nLocation: Remote\nPython Django APIs',
		))
		Applicant.objects.create(
			first_name='Taylor',
			middle_name='',
			last_name='Rivera',
			contact_number='+63 912 222 3344',
			email='taylor@example.com',
			position=job.title,
			applied_job=job,
			job_preference_locked=True,
			match_score=88,
			match_notes={
				'overall_score': 40,
				'target_job_id': job.id,
				'scored_job_ids': [job.id],
				'matching_skills': ['Python', 'Django'],
				'missing_skills': ['Kubernetes'],
				'recommendation': {
					'job_id': 999,
					'job_title': 'Kids Attendant',
					'overall_score': 73,
				},
				'breakdown': {
					'skill_score': 10,
					'experience_score': 20,
					'education_score': 10,
				},
			},
			resume=build_docx_file(
				'taylor.docx',
				[
					'Taylor Rivera',
					'taylor@example.com',
					'+63 912 222 3344',
				],
			),
		)

		response = self.client.get(reverse('dashboard'))

		self.assertEqual(response.status_code, 200)
		self.assertContains(response, 'Analyze Score')
		self.assertContains(response, 'scoreBreakdownDialog-dashboard-', html=False)
		self.assertContains(response, 'Taylor Rivera Score Analysis')
		self.assertContains(response, 'Skill Match')
		self.assertContains(response, '10%')
		self.assertContains(response, 'class="db-score-bar-track"', html=False)
		self.assertContains(response, 'class="db-score-bar-fill" data-score="10"', html=False)
		self.assertContains(response, '40%')
		self.assertContains(response, 'Matching Skills')
		self.assertContains(response, 'Missing Skills')
		self.assertContains(response, 'Kubernetes')
		self.assertContains(response, 'Recommendation')
		self.assertContains(response, 'Kids Attendant')
		self.assertContains(response, '73%')

	def test_dashboard_removes_applicant_from_status_tab(self):
		applicant = Applicant.objects.create(
			first_name='Delete',
			last_name='Me',
			contact_number='+63 912 345 6789',
			email='delete-me@example.com',
			resume=build_docx_file(
				'delete-me.docx',
				[
					'Delete Me',
					'delete-me@example.com',
					'+63 912 345 6789',
				],
			),
		)

		response = self.client.post(
			reverse('dashboard'),
			{'action': 'remove_applicant', 'applicant_id': str(applicant.id), 'redirect_panel': 'interview'},
			follow=True,
		)

		self.assertEqual(response.status_code, 200)
		self.assertFalse(Applicant.objects.filter(id=applicant.id).exists())
		self.assertContains(response, 'Delete Me was removed.')

	def test_logout_clears_dashboard_session(self):
		response = self.client.post(reverse('logout'), follow=True)

		self.assertEqual(response.status_code, 200)
		self.assertFalse('_auth_user_id' in self.client.session)
		self.assertContains(response, 'Logged out successfully.')


class ResumeAccessTests(TestCase):
	def setUp(self):
		self.staff_user = User.objects.create_user(username='hr.staff@example.com', password='secret123', is_staff=True)
		self.applicant = Applicant.objects.create(
			first_name='Secure',
			last_name='Applicant',
			contact_number='+63 912 345 6789',
			email='secure-applicant@example.com',
			resume=build_docx_file(
				'secure-applicant.docx',
				[
					{'style': 'Heading 1', 'runs': [('Secure Applicant', True)]},
					'secure-applicant@example.com',
					'+63 912 345 6789',
					('Python and Django development', 'List Bullet'),
				],
			),
			position='Backend Developer',
			resume_text='Secure Applicant\nPython\nDjango\nREST API',
			match_score=78,
			match_notes={
				'overall_score': 78,
				'matching_skills': ['Python'],
				'missing_skills': ['Django'],
				'recommendations': [
					{'job_id': 1, 'job_title': 'Backend Developer', 'overall_score': 78},
					{'job_id': 2, 'job_title': 'API Engineer', 'overall_score': 72},
				],
				'breakdown': {
					'skill_score': 65,
					'experience_score': 80,
					'education_score': 90,
				},
			},
		)

	def test_dashboard_uses_protected_resume_link(self):
		self.client.force_login(self.staff_user)

		response = self.client.get(reverse('dashboard'))

		protected_url = reverse('serve_resume', args=[Path(self.applicant.resume.name).name])
		self.assertEqual(response.status_code, 200)
		self.assertContains(response, protected_url)
		self.assertNotContains(response, self.applicant.resume.url)

	def test_protected_resume_download_requires_staff_login(self):
		response = self.client.get(reverse('serve_resume', args=[Path(self.applicant.resume.name).name]))

		self.assertEqual(response.status_code, 302)
		self.assertIn(reverse('login'), response['Location'])

	def test_applicant_detail_requires_staff_login(self):
		response = self.client.get(reverse('applicant_detail', args=[self.applicant.id]))

		self.assertEqual(response.status_code, 302)
		self.assertIn(reverse('login'), response['Location'])

	def test_staff_can_view_applicant_detail_with_resume_and_score_analysis(self):
		self.client.force_login(self.staff_user)

		response = self.client.get(reverse('applicant_detail', args=[self.applicant.id]))

		self.assertEqual(response.status_code, 200)
		self.assertContains(response, self.applicant.get_full_name())
		self.assertNotContains(response, 'Open Resume')
		self.assertContains(response, reverse('applicant_resume_preview', args=[self.applicant.id]))
		self.assertContains(response, 'Score Analysis')
		self.assertContains(response, 'Extracted Resume Text')
		self.assertContains(response, 'All Job Position Scores')
		self.assertContains(response, 'Python')
		self.assertContains(response, 'Django')
		self.assertContains(response, '78%')
		self.assertContains(response, 'REST API')
		self.assertContains(response, 'Backend Developer')
		self.assertContains(response, 'API Engineer')
		self.assertContains(response, '72%')
		self.assertContains(response, 'Contact Applicant')
		self.assertContains(response, 'https://mail.google.com/mail/?', html=False)
		self.assertContains(response, 'sms:+639120001111', html=False)
		self.assertContains(response, 'Send Email')
		self.assertContains(response, 'Send SMS')
		self.assertEqual(len(response.context['job_position_scores']), 1)
		self.assertEqual(response.context['job_position_scores'][0]['job_title'], 'API Engineer')
		self.assertEqual(response.context['job_position_scores'][0]['overall_score'], 72)

	def test_applicant_resume_preview_requires_staff_login(self):
		response = self.client.get(reverse('applicant_resume_preview', args=[self.applicant.id]))

		self.assertEqual(response.status_code, 302)
		self.assertIn(reverse('login'), response['Location'])

	def test_staff_can_view_pdf_style_resume_preview_for_docx(self):
		self.client.force_login(self.staff_user)

		response = self.client.get(reverse('applicant_resume_preview', args=[self.applicant.id]))

		self.assertEqual(response.status_code, 200)
		self.assertEqual(response['Content-Type'], 'text/html; charset=utf-8')
		self.assertNotIn('X-Frame-Options', response)
		self.assertContains(response, 'Resume Preview')
		self.assertContains(response, 'secure-applicant.docx')
		self.assertContains(response, 'docx-preview.min.js')
		self.assertContains(response, reverse('serve_resume', args=[Path(self.applicant.resume.name).name]))
		self.assertContains(response, 'Loading DOCX preview...')
		self.assertContains(response, 'docxPreviewContainer')
		self.assertNotContains(response, '<pre class="preview-text">', html=False)

	def test_staff_can_view_resume_from_protected_route(self):
		self.client.force_login(self.staff_user)

		response = self.client.get(reverse('serve_resume', args=[Path(self.applicant.resume.name).name]))

		self.assertEqual(response.status_code, 200)
		self.assertEqual(
			response['Content-Type'],
			'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
		)
		self.assertIn('inline', response['Content-Disposition'])
		self.assertIn(Path(self.applicant.resume.name).name, response['Content-Disposition'])

	def test_direct_media_resume_url_is_not_served(self):
		response = self.client.get(self.applicant.resume.url)

		self.assertEqual(response.status_code, 404)

	def test_sensitive_applicant_data_is_encrypted_at_rest(self):
		self.applicant.resume_text = 'Python resume text with private details'
		self.applicant.parsed_skills = ['Python', 'Django']
		self.applicant.parsed_education = ['Bachelor of Science in Computer Science']
		self.applicant.match_notes = {'overall_score': 91, 'matching_skills': ['Python']}
		self.applicant.save(update_fields=['resume_text', 'parsed_skills', 'parsed_education', 'match_notes'])

		with connection.cursor() as cursor:
			cursor.execute(
				'SELECT first_name, email, resume, resume_text, parsed_skills, match_notes FROM main_applicant WHERE id = %s',
				[self.applicant.id],
			)
			row = cursor.fetchone()

		self.assertIsNotNone(row)
		self.assertNotEqual(row[0], 'Secure')
		self.assertNotEqual(row[1], 'secure-applicant@example.com')
		self.assertNotEqual(row[3], 'Python resume text with private details')
		self.assertNotEqual(row[4], '["Python", "Django"]')
		self.assertNotEqual(row[5], '{"overall_score": 91, "matching_skills": ["Python"]}')
		self.assertTrue(str(row[2]).startswith('gAAAA'))

		with open(self.applicant.resume.path, 'rb') as stored_file:
			stored_bytes = stored_file.read()

		self.assertTrue(stored_bytes.startswith(b'gAAAA'))


class AuthViewTests(TestCase):
	def test_login_shows_pending_approval_message_for_inactive_hr_user(self):
		User.objects.create_user(
			username='pending.hr@example.com',
			email='pending.hr@example.com',
			password='secret123',
			is_staff=True,
			is_active=False,
		)

		response = self.client.post(
			reverse('login'),
			{
				'username': 'pending.hr@example.com',
				'password': 'secret123',
			},
		)

		self.assertEqual(response.status_code, 200)
		self.assertContains(response, 'Wait for Admin approval.')
		self.assertNotContains(response, 'Wrong Email or Password. Wait for Admin approval.')

	def test_login_keeps_generic_message_for_invalid_credentials(self):
		response = self.client.post(
			reverse('login'),
			{
				'username': 'missing@example.com',
				'password': 'bad-password',
			},
		)

		self.assertEqual(response.status_code, 200)
		self.assertContains(response, 'Wrong Email or Password')
		self.assertNotContains(response, 'Wait for Admin approval.')